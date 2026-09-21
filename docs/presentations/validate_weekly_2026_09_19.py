"""Check the delivered deck, embedded media, rendered text, and browser package."""
from pathlib import Path
from zipfile import ZipFile
from urllib.parse import unquote
import hashlib
import json
import posixpath
import re
import sys
import unicodedata
from lxml import etree

if Path('/tmp/topogym_pptx_deps').exists():
    sys.path.insert(0, '/tmp/topogym_pptx_deps')
from pptx import Presentation

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
NAME = 'TopoGym_Weekly_2026-09-19'
ASSETS = OUT / '.weekly_2026_09_19'
manifest = json.loads((OUT / f'{NAME}.manifest.json').read_text())
media = json.loads((ASSETS / 'embedded_media.json').read_text())
archived = {v['path']: v for v in json.loads((ROOT / 'reports/weekly/2026-09-19/video_validation.json').read_text())['videos']}
prs = Presentation(OUT / f'{NAME}.pptx')
ns = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'p14':'http://schemas.microsoft.com/office/powerpoint/2010/main'}
verified = []
with ZipFile(OUT / f'{NAME}.pptx') as z:
    assert z.testzip() is None
    parts = [p for p in z.namelist() if p.startswith('ppt/media/') and p.endswith('.mp4')]
    assert len(parts) == 39
    ct = etree.fromstring(z.read('[Content_Types].xml'))
    assert any(e.get('Extension') == 'mp4' and e.get('ContentType') == 'video/mp4' for e in ct)
    for m in manifest['embedded_movies']:
        n = m['slide']
        slide = etree.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
        rels = {e.get('Id'): e for e in etree.fromstring(z.read(f'ppt/slides/_rels/slide{n}.xml.rels'))}
        pic = next(p for p in slide.findall('.//p:pic', ns)
                   if p.find('p:nvPicPr/p:cNvPr', ns).get('id') == str(m['shape_id']))
        link = pic.find('p:nvPicPr/p:cNvPr/a:hlinkClick', ns)
        assert link.get('action') == 'ppaction://media' and not link.get('{'+ns['r']+'}id')
        embed = pic.find('.//p14:media', ns).get('{'+ns['r']+'}embed')
        rel = rels[embed]
        assert rel.get('TargetMode') != 'External'
        part = posixpath.normpath(posixpath.join('ppt/slides', rel.get('Target')))
        assert hashlib.sha256(z.read(part)).hexdigest() == m['sha256']
        video_rel = rels[pic.find('.//a:videoFile', ns).get('{'+ns['r']+'}link')]
        assert video_rel.get('Target') == rel.get('Target') and video_rel.get('TargetMode') != 'External'
        targets = slide.findall('.//p:timing//p:video/p:cMediaNode/p:tgtEl/p:spTgt', ns)
        assert any(t.get('spid') == str(m['shape_id']) for t in targets)
        assert hashlib.sha256(Path(m['source']).read_bytes()).hexdigest() == archived[m['source']]['sha256']
        verified.append(dict(slide=n, group=m['group'], camera=m['camera'], part=part,
                             internal_media=True, media_click_action=True, sha256=m['sha256']))

for h in manifest['hyperlinks']:
    assert (OUT / unquote(h['target'])).resolve().is_file(), h
for s in prs.slides:
    for shape in s.shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= prs.slide_width + 2
        assert shape.top + shape.height <= prs.slide_height + 2

pdf = etree.parse('/tmp/topogym_weekly_slides_bbox.xml')
pns = {'x':'http://www.w3.org/1999/xhtml'}
pages = pdf.findall('.//x:page', pns)
assert len(pages) == len(prs.slides) == 20
def norm(t):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', t))
mismatches = []
text_boxes = 0
for i, slide in enumerate(prs.slides):
    words = [(float(w.get('xMin')), float(w.get('yMin')), float(w.get('xMax')), float(w.get('yMax')), w.text or '')
             for w in pages[i].findall('.//x:word', pns)]
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text.strip():
            continue
        text_boxes += 1
        x, y, w, h = [v/12700 for v in [shape.left, shape.top, shape.width, shape.height]]
        found = [v for v in words if x-.5 <= (v[0]+v[2])/2 <= x+w+.5 and y-.5 <= (v[1]+v[3])/2 <= y+h+.5]
        actual = ''.join(v[4] for v in sorted(found, key=lambda v:(round(v[1]/3),v[0])))
        if norm(shape.text) not in norm(actual):
            mismatches.append(dict(slide=i+1, expected=shape.text, actual=actual))

web = json.loads((ASSETS / 'web_slides.json').read_text())
with ZipFile(OUT / f'{NAME}_browser.zip') as z:
    assert z.testzip() is None
    members = set(z.namelist())
    assert f'{NAME}.html' in members
    count = 0
    for s in web:
        assert unquote(s['background']) in members
        for m in s['movies']:
            assert unquote(m['src']) in members and unquote(m['poster']) in members
            count += 1
    assert count == 39
assert {m['group'] for m in web[8]['movies']} == {'claws_a'}
assert all(m['duration'] == 158.6 for m in web[8]['movies'])
coverage = {m['source'] for m in manifest['embedded_movies']} | {h['source'] for h in manifest['hyperlinks'] if h['source'].endswith('.mp4')}
assert len(coverage) == 73
result = dict(slides=20, rendered_pdf_pages=20, all_shapes_within_slide=True,
              all_text_present_in_render=not mismatches, text_boxes_checked=text_boxes,
              text_mismatches=mismatches, text_validation_method='PDF words selected by each original textbox rectangle',
              embedded_video_count=39, embedded_media=verified,
              all_embedded_media_hashes_match=True, all_original_archive_hashes_unchanged=True,
              all_playback_copies_fully_decoded=all(m['full_decode'] for m in media),
              all_video_frames_fps_and_durations_preserved=all(m['frames_and_timing_preserved'] for m in media),
              hyperlinks_checked=len(manifest['hyperlinks']), unique_mp4_links=manifest['linked_mp4_unique'],
              unique_demo_coverage=73, claws_primary_variant='A0', browser_package_media_count=39,
              powerpoint_desktop_playback_tested=False,
              files={})
browser_check = ASSETS / 'browser_playback_validation.json'
if browser_check.exists():
    result['browser_playback_verification'] = json.loads(browser_check.read_text())
for ext in ['pptx', 'pdf', 'html', '_browser.zip']:
    p = OUT / (NAME + ('' if ext.startswith('_') else '.') + ext)
    result['files'][p.name] = dict(bytes=p.stat().st_size, sha256=hashlib.sha256(p.read_bytes()).hexdigest())
(OUT / f'{NAME}.validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
print(json.dumps({k:v for k,v in result.items() if k not in ['embedded_media','files','browser_playback_verification']}, ensure_ascii=False, indent=2))
assert not mismatches, 'Rendered text mismatch; inspect validation.json'
