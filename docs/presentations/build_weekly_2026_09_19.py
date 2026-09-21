"""Build the editable presentation from the frozen 2026-09-19 weekly report.

Requires python-pptx and Pillow. Embeds 39 complete demonstration videos
and synchronized poster frames; supplemental originals retain relative links.
"""
from pathlib import Path
import json
import os
import sys
from urllib.parse import quote

if Path('/tmp/topogym_pptx_deps').exists():
    sys.path.insert(0, '/tmp/topogym_pptx_deps')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DATA = ROOT / 'reports/weekly/2026-09-19'
ASSETS = OUT / '.weekly_2026_09_19'
DECK = OUT / 'TopoGym_Weekly_2026-09-19.pptx'
report = json.loads((DATA / 'report.json').read_text())
inventory = json.loads((DATA / 'demo_inventory.json').read_text())
groups = {g['id']: g for g in inventory}
audit = json.loads((DATA / 'token_audit.json').read_text())
videos = {v['path']: v for v in json.loads((DATA / 'video_validation.json').read_text())['videos']}
thumbs = {(t['group'], t['camera']): t for t in json.loads((ASSETS / 'thumbnails.json').read_text())}
media = {(t['group'], t['camera']): t for t in json.loads((ASSETS / 'embedded_media.json').read_text())}

FONT = 'Noto Sans CJK SC'
C = dict(bg='F5F7FA', white='FFFFFF', ink='14283F', muted='65748A',
         teal='137D69', teal_light='E0F1EB', amber='A56C18', amber_light='FFF0CE',
         red='AD4D48', red_light='F7E6E4', line='DBE2EB', grey='E7EBF1', dark='101E31')
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'TopoGym 周报 | 2026-09-19'
prs.core_properties.subject = '任务进展、GPT-6 Astra 2M 目标预算测试、reviewer 讨论'
prs.core_properties.author = 'TopoGym'
prs.core_properties.keywords = 'TopoGym, robotics, weekly report, GPT-6 Astra'
slide_records = []
links = []
embedded_movies = []


def rgb(s):
    return RGBColor.from_string(C.get(s, s))


def rect(slide, x, y, w, h, fill='white', border=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = 0.09
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if border:
        shape.line.color.rgb = rgb(border)
        shape.line.width = Pt(0.6)
    else:
        shape.line.fill.background()
    return shape


def set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    props = run._r.get_or_add_rPr()
    for tag in ['a:ea', 'a:cs']:
        el = OxmlElement(tag)
        el.set('typeface', FONT)
        props.append(el)


def text(slide, s, x, y, w, h, size=18, color='ink', bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.015)
    tf.margin_top = tf.margin_bottom = Inches(0.015)
    tf.vertical_anchor = valign
    for i, line in enumerate(s.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.line_spacing = 1.14
        run = p.add_run()
        run.text = line
        set_font(run, size, color, bold)
    return shape


def target(path):
    p = Path(path)
    if not p.is_absolute():
        p = ROOT / p
    assert p.is_file(), p
    return quote(os.path.relpath(p, OUT), safe='/._-')


def hyperlink(shape, path, label=None):
    shape.click_action.hyperlink.address = target(path)
    links.append({'slide': len(prs.slides), 'label': label or shape.name,
                  'source': str(path), 'target': target(path)})


def linked_text(slide, s, path, x, y, w, h=.27, size=11.5, color='teal'):
    shape = text(slide, s, x, y, w, h, size, color)
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.hyperlink.address = target(path)
            r.font.underline = True
    links.append({'slide': len(prs.slides), 'label': s, 'source': str(path), 'target': target(path)})
    return shape


def note(slide, s, sources=()):
    tail = '\n\n来源（相对仓库根目录）：\n' + '\n'.join(str(x) for x in sources)
    slide.notes_slide.notes_text_frame.text = s + tail


def base(title, section, subtitle='', dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb('dark' if dark else 'bg')
    fg = 'white' if dark else 'ink'
    rect(slide, .46, .30, .08, .20, 'teal')
    text(slide, section, .64, .27, 11.6, .3, 11, 'teal' if not dark else '77CDB8', True)
    text(slide, title, .46, .70, 12.35, .59, 29, fg, True)
    if subtitle:
        text(slide, subtitle, .48, 1.34, 12.25, .38, 13, 'muted' if not dark else 'B9C6D6')
    rect(slide, .48, 7.12, 12.35, .008, 'line' if not dark else '33435A')
    text(slide, 'TopoGym  /  2026.09.19  /  冻结周报快照', .48, 7.20, 10.7, .19, 8.5,
         'muted' if not dark else 'A0AFC2')
    text(slide, f'{len(prs.slides):02d}', 12.05, 7.17, .76, .26, 10, fg, align=PP_ALIGN.RIGHT)
    slide_records.append({'number': len(prs.slides), 'title': title, 'section': section})
    return slide


def picture_fit(slide, path, x, y, w, h):
    im = Image.open(path)
    iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w-pw)/2), Inches(y + (h-ph)/2),
                                    width=Inches(pw), height=Inches(ph))


camera_names = {'agentview': '全景', 'front_oblique': '斜前', 'wrist': '腕部',
                'cam_high': '高位', 'cam_low': '低位', 'cam_left_wrist': '左腕',
                'cam_right_wrist': '右腕', 'detail': '近景'}


def preview(slide, gid, cam, x, y, w=3.78, h=1.65):
    row = thumbs[(gid, cam)]
    video = media[(gid, cam)]
    rect(slide, x, y, w, h, '182438', radius=True)
    iw, ih = Image.open(row['image']).size
    scale = min((w-.05)/iw, (h-.05)/ih)
    pw, ph = iw*scale, ih*scale
    movie = slide.shapes.add_movie(video['playback_file'], Inches(x+(w-pw)/2),
                                   Inches(y+(h-ph)/2), Inches(pw), Inches(ph),
                                   poster_frame_image=row['image'], mime_type='video/mp4')
    movie.name = f'DEMO {gid}/{cam}'
    # Keep the media click action from add_movie; an external hyperlink would
    # override in-slide playback and recreate the original delivery problem.
    duration = float(video['playback_probe']['duration'])
    caption = f'▶ {camera_names[cam]} · 点击播放 · {duration:.1f} s'
    text(slide, caption, x+.02, y+h+.035, w-.04, .27, 11.5, 'teal')
    embedded_movies.append(dict(slide=len(prs.slides), shape_id=movie.shape_id,
                                group=gid, camera=cam, source=row['source'],
                                playback_file=video['playback_file'], sha256=video['sha256'],
                                duration_s=duration, x=x+(w-pw)/2, y=y+(h-ph)/2,
                                width=pw, height=ph, poster=row['image']))


def demo_row(slide, gid, label, y):
    g = groups[gid]
    text(slide, label, .49, y, 2.53, .28, 11.1, 'muted')
    x = 3.01
    for cam, path in g['videos'].items():
        linked_text(slide, camera_names[cam], path, x, y, .80, .28, 11.1)
        x += .87


def info_block(slide, label, body, y, color='teal', h=1.30):
    rect(slide, 8.38, y, 4.46, h, 'white', radius=True)
    rect(slide, 8.38, y, .045, h, color)
    text(slide, label, 8.57, y+.12, 4.03, .28, 12, color, True)
    text(slide, body, 8.57, y+.48, 4.01, h-.52, 16, 'ink')


def chip(slide, label, x, y, w, fill='teal_light', color='teal'):
    rect(slide, x, y, w, .36, fill, radius=True)
    text(slide, label, x+.05, y+.04, w-.1, .26, 10.5, color, True, PP_ALIGN.CENTER)


# 01 — cover and agenda
s = base('TopoGym 研究周报', 'WEEKLY RESEARCH UPDATE', dark=True)
text(s, '从可运行场景，走向可验证的变体数据集', .48, 1.46, 11.6, .54, 25, 'white')
text(s, '工程进展 18:27 CDT  ·  Agent 记录 18:24:55 CDT', .49, 2.11, 11.6, .33, 13, 'B9C6D6')
text(s, '视频观看：进入 PowerPoint 放映模式后点击画面；39 路完整视频已内嵌，可离线播放。', .49, 2.63, 12.0, .29, 11.8, 'B9C6D6')
for j, (num, name, body) in enumerate([
    ('01', '11 项任务', '场景、物理 demo、随机生成与逐例验收'),
    ('02', 'GPT-6 Astra', '2M 目标预算下的尝试、结果与计量核对'),
    ('03', 'Reviewer 问题', '可解性证据与直接 session 测试方法')]):
    x = .49 + j*4.19
    rect(s, x, 3.08, 4.02, 2.05, '1B2D44', radius=True)
    text(s, num, x+.22, 3.23, 1.0, .45, 29, '77CDB8', True)
    text(s, name, x+.22, 3.86, 3.56, .4, 22, 'white', True)
    text(s, body, x+.22, 4.43, 3.56, .56, 14, 'B9C6D6')
text(s, '核心结论', .50, 5.64, 1.5, .3, 12, '77CDB8', True)
text(s, '生成器已有明显进展；逐变体物理资格与严格 2M 评测仍需补齐。', .50, 6.05, 12.1, .56, 22, 'white', True)
note(s, '本演示文稿转换自已完成的周报，不更新正在运行的工程或求解结果。三部分与 Markdown 周报一致。',
     ['reports/weekly/2026-09-19/weekly-report.md', 'reports/weekly/2026-09-19/sources.json'])

# 02 — editable milestone matrix
s = base('11 项任务：四阶段进展总览', '01 / 任务进展', '按明确实例与协议范围标记；色块不表示工时百分比或模型成功率。')
xs = [4.18, 6.36, 8.54, 10.72]
for j, h in enumerate(['① 场景生成', '② 物理＋完整 demo', '③ 随机生成接口', '④ 逐变体稳定可解']):
    text(s, h, xs[j], 1.93, 2.07, .4, 13, 'ink', True, PP_ALIGN.CENTER)
status = {2: ('teal_light', 'teal', '已达当前范围'), 1: ('amber_light', 'amber', '部分 / 验证中'),
          0: ('grey', 'muted', '未开始')}
for i, task in enumerate(report['tasks']):
    y = 2.48 + i*.34
    text(s, task['name'], .53, y+.005, 3.63, .30, 12.3)
    for j, val in enumerate(task['stages']):
        bg, fg, label = status[val]
        rect(s, xs[j], y, 2.05, .29, bg, radius=True)
        text(s, label, xs[j]+.04, y+.025, 1.97, .24, 10.6, fg, False, PP_ALIGN.CENTER)
text(s, '注：Claws 的物理通过对应 B0/C0；Blanket 按当前麻花初态计。所有任务均未完成全量变体验收。', .53, 6.48, 12.1, .41, 12.1, 'muted')
note(s, '绿色不取消预抓取、辅助抓持或末态仍持稳等协议边界。Fiber 严格线形收敛和 Knot 完整物理资格尚未通过。',
     ['reports/weekly/2026-09-19/report.json', 'reports/weekly/2026-09-19/weekly-report.md'])

task_slides = [
 dict(id='cast_coil', title='Cast Coil｜卷绕件装配', gid='coil_base', subtitle='固定三家族 + 条件微调已接入 seed；随机动力学验收尚未闭环。',
      done='384 个不同配置可复现\n48/48 几何；12/12 双臂规划',
      gap='六次代表执行均完成装配释放，严格归档仅 4/6 通过。',
      next='保留 A/C 的微米级接触失败；对随机实例增加物理验收。',
      scope='基础版 33.15 s｜预抓取＋接触辅助；装配、放置、释放、静置。',
      more=[('coil_base','基础装配 · 五路')], source='docs/develop/tasks/cast_coil_seed_generator_zh.md'),
 dict(id='cast_keyhole', title='Cast Keyhole｜钥匙孔拆解', gid='keyhole_base', subtitle='F1 长行程穿齿、F2 偏心绕行、F3 转向后回退；三族各三种外形。',
      done='256 seed 复现；200 次实体/边界检查；3 实例 × 2 步长通过。',
      gap='物理资格止于预抓取、辅助拆解后持稳；不含放桌松手。',
      next='每个新实例独立验收；生成端与 agent 终态协议继续对齐。',
      scope='基础版 23.92 s｜完整分离后双手仍保持抓持；实测状态重渲染。',
      more=[('keyhole_base','基础拆解 · 五路')],source='docs/develop/tasks/cast_keyhole_variants_zh.md'),
 dict(id='horseshoe_ring', title='Horseshoe Ring｜马蹄环', gid='horseshoe_base', subtitle='高位 / 外侧 / 后撤接近形成三族；路线差异约 20–22 cm。',
      done='256 seed 参数检查；48 项初态/折合 IK；逐实例验收入口已接通。',
      gap='新族只完成 2 ms 验收，未做 1 ms 复验；不能继承基础版资格。',
      next='最终条件域放行 3 个实例；补齐新族半步长和稳定性验收。',
      scope='基础版 113.70 s｜辅助抓持、交接、取环与自由落桌；2/1 ms 复验通过。',
      more=[('horseshoe_base','重力基础版 · 五路')],source='docs/develop/tasks/horseshoe_ring_variant_generator_zh.md'),
 dict(id='hedgehog_in_cage', title='Hedgehog in Cage｜刺球出笼', gid='hedgehog_side', subtitle='当前默认池：单斜窗、固定六刺侧孔、GEN 顶部钥匙孔。',
      done='两种新原型完整取出/释放，两档子步通过；13 实例预检通过。',
      gap='随机后代尚未物理执行。侧孔与 CLASSIC 共享侧向取出策略。',
      next='沿用 v2 seed 与重载接口，逐例补动力学；不混入 v3 几何预演。',
      scope='侧孔 v2 · 82.25 s｜预抓取＋接触辅助；取出、落桌、释放、静置。', metric=('3','当前家族','1000','seed 家族抽样'),
      more=[('hedgehog_classic','CLASSIC · 基础测试'),('hedgehog_oblique','单斜窗 · 当前家族 A'),('hedgehog_gen','GEN · 当前家族 C')],source='docs/develop/tasks/hedgehog_in_cage_zh.md'),
 dict(id='lock_picking', title='Lock Picking｜开锁', gid='lock_base', subtitle='F1/F2/F3 必要回退 1/2/3 次：体现序列复杂度，尚无实测难度排序。',
      done='256 seed 在最多 8 次尝试内通过几何/IK；九个 S/M/L 代表通过。',
      gap='99 个 seed 需要重试；新三族完整抓起、动力学和松手未验收。',
      next='将已对齐的 IK 路径接到完整抓取与释放流程，再逐例验收。',
      scope='基础版 159.35 s｜成功物理状态重渲染；被动夹爪连杆按开合量重建。', metric=('256','几何 / IK 接受','99','需要确定性重试'),
      more=[('lock_base','两回退基础版 · 三路')],source='docs/develop/tasks/lock_picking_seed_generator_zh.md'),
 dict(id='fiber_bundling', title='Fiber Bundling｜光纤收束', gid='fiber_base', subtitle='保留稀疏 / 标准 / 密集三类；按已确认规则，不要求三种本质不同 IK。',
      done='300 seed 几何、39 边界组合通过；3 随机实例短静置通过。',
      gap='基础 8 芯两步长均穿环成功，但严格全程线形收敛未通过。',
      next='按实例检查 IK、细线禁碰、完整穿环和数值稳定性。',
      scope='8 芯 / R13.5 mm · 177 s｜只抓粗护套、接触后辅助；末态保持抓持。', metric=('300','不同参数实例','39','N/R 边界组合'),
      more=[('fiber_base','基础无 UI 视频 · 三路')],source='docs/develop/tasks/fiber_bundling_variant_pipeline_zh.md'),
 dict(id='interlocked_claws', title='Interlocked Claws｜A0 基础变体', gid='claws_a', subtitle='本页播放 A0 原版基础 demo；B0/C0 的新增物理验收另列，不替代 A0 资格。',
      done='A0 五路基础 demo 已归档；B0/C0 新原型通过完整物理复验。',
      gap='A0 使用理想预抓取辅助；Genesis 非凸接触资格尚未通过。',
      next='发布新版本生成器，重新配对实例、IK 和物理资格。',
      scope='A0 基础版 · 158.6 s｜理想预抓取辅助全过程；非凸接触资格限制保留。',
      more=[('claws_a','A0 基础版 · 原始五路'),('claws_b','B0 螺旋冠 · 补充五路'),('claws_c','C0 正交双窗 · 补充五路')],source='docs/develop/tasks/interlocked_claws_zh.md'),
 dict(id='knot_untying', title='Knot Untying｜解绳结', gid='knot_base', subtitle='三叶结已有连续轨迹；八字结、双重单结已建模，连续执行仍在验证。',
      done='两种新结型各有 10 s 独立静置；18 组收窄边界首段预检通过。',
      gap='v9 的 P3 网格收敛及通用双手交接未过；新族尚无完整成功。',
      next='完成独立解结路线与释放验收；easy/medium/hard 仍是目标档。',
      scope='三叶结 v9 · 75 s｜接触后辅助，连续解结、释放；完整物理套件未通过。',
      more=[('knot_base','三叶结 v9 · 五路')],source='docs/develop/tasks/knot_untying_family_space_zh.md'),
 dict(id='buzz_wire_3d', title='Buzz Wire 3D｜立体电流棒', gid='buzz_f3', subtitle='F3 四关 v3 主运行通过；F2 后半程仍在搜索。',
      done='F3 的 10001 点密集 IK 通过；106.92 s 真实摩擦抓持主执行完成。',
      gap='F3 独立 2 ms 复验重跑中；随机范围尚未逐实例认证。',
      next='完成 F3 复验与归档，继续 F2 完整路线及机器人净空检查。',
      scope='F3 成片 106.96 s｜预穿环/预抓取、纯摩擦抓持、释放滑落；实测状态重渲染。', metric=('10001','密集 IK 检查点','106.92 s','主物理执行完成'),
      more=[('buzz_f3','F3 最新主运行 · 三路'),('buzz_base','紧凑基础版 · 三路')],source='runs/environment-solvability/buzz_wire_3d/family_solvability_20260919/result.md'),
 dict(id='rope_threading', title='Rope Threading｜绳穿三环', gid='rope_prefix', subtitle='最新可播放录像是局部前段；三环完整参考执行仍在进行。',
      done='抓取、抬升、释放与静置前段通过；五路同步录像共 29.05 s。',
      gap='全流程快照 19.005 s，双手抓持有效；尚未获得三环完整成功。',
      next='继续长程执行；解决加载半步长迟发运动与静置验收。',
      scope='局部前段 29.05 s｜包含脚本主动释放；不等于完成三环任务。',
      more=[('rope_prefix','前段验证 · 五路')],source='docs/develop/tasks/rope_threading_zh.md'),
 dict(id='blanket_flatten', title='Blanket Flatten｜被子展平', gid='blanket_latest', subtitle='当前选定麻花初态稳定；尚未获得复杂初态的成功展平 demo。',
      done='独立静置 10 s 通过，最大漂移 0.143 mm；v10 三路已完整录制。',
      gap='v10 八次保护释放，末态覆盖 45.4%；v8 为 75.7%，仍未达标。',
      next='解决解扭、被覆盖布角与拉伸控制；旧简单初态成功单独保留。',
      scope='最新 v10 · 179.5 s / 2 fps｜未完成诊断；局部粒子附着辅助，实测状态重渲染。', metric=('45.4%','v10 最新末态覆盖','75.7%','v8 较好试验末态'),
      more=[('blanket_latest','v10 最新未完成 · 三路'),('blanket_best','v8 补充进展 · 三路'),('blanket_old','旧简单初态成功 · 三路')],source='runs/environment_bootstrap/blanket_flatten/gathered_robot_study/result.md'),
]

for d in task_slides:
    s = base(d['title'], '01 / 任务进展', d['subtitle'])
    gid = d['gid']
    cams = ['agentview', 'front_oblique', 'wrist'] if 'agentview' in groups[gid]['videos'] else ['cam_high','cam_low','cam_left_wrist','cam_right_wrist']
    for i, cam in enumerate(cams):
        preview(s, gid, cam, .49+(i%2)*3.94, 1.94+(i//2)*2.03)
    if len(cams)==3:
        a, al, b, bl = d['metric']
        rect(s, 4.43, 3.97, 3.78, 1.65, 'white', radius=True)
        text(s, a, 4.64, 4.12, 3.25, .41, 25, 'teal', True)
        text(s, al, 4.64, 4.61, 3.25, .30, 12, 'muted')
        text(s, b, 4.64, 4.98, 1.7, .42, 21, 'ink', True)
        text(s, bl, 6.26, 5.07, 1.72, .34, 10.6, 'muted')
    info_block(s, '本周已完成', d['done'], 1.94, h=1.37)
    info_block(s, '当前边界', d['gap'], 3.46, color='amber', h=1.40)
    info_block(s, '下一步', d['next'], 5.01, h=1.40)
    text(s, d['scope'], .49, 1.73, 7.72, .22, 9.6, 'muted')
    for j, (g, label) in enumerate(d['more']):
        demo_row(s, g, label, 6.00+j*.32)
    linked_text(s, '验证记录 ↗', d['source'], 8.55, 6.63, 3.5, .27, 11.8)
    source_groups = [g for g in inventory if g['task_id']==d['id']]
    detailed = '\n\n'.join(f"{g['title']}\n{g['scope']}\n"+'\n'.join(f"{c}: {p}" for c,p in g['videos'].items()) for g in source_groups)
    note(s, d['subtitle']+'\n'+d['done']+'\n'+d['gap']+'\n'+d['next']+'\n\n'+detailed,
         [d['source'], 'reports/weekly/2026-09-19/demo_inventory.json'])

# 14 — actual outcomes, separated from budget compliance
s = base('6 个任务、7 次运行：尚无成功求解', '02 / GPT-6 ASTRA 测试', '6 次已结束为 failed；Claws 在快照时仍为 ready。工程 demo 与自主求解分开统计。')
rows = [
 ('Coil 001','11 / 8.30 s','主动停止','能抓放银件，未建立装配办法'),
 ('Coil 002','64 / 74.34 s','主动停止','右夹爪持续不响应张开，根因未确认'),
 ('Horseshoe','2 / 1.64 s','预算终止','抬转左马蹄，尚未取下圆环'),
 ('Keyhole','19 / 9.01 s','主动停止','内部槽齿附近受阻，未找到完整分离路径'),
 ('Hedgehog','7 / 10.05 s','主动停止','部分移出，仍受笼柱阻挡，未取出/释放'),
 ('Lock Picking','10 / 16.29 s','自动 failed','推进第一锁片时，机构接触力 >10 N'),
 ('Claws','14 / 9.86 s','进行中 / ready','仍然互锁；不预判最终成败'),
]
col=[.51, 2.58, 4.67, 6.67]
for j, t in enumerate(['任务 / run', '段数 / 控制时间', '终止类型', '主要结果']):
    text(s,t,col[j],1.98,[1.96,1.96,1.83,5.65][j],.30,13,'muted',True)
for i,row in enumerate(rows):
    y=2.49+i*.48
    rect(s,.48,y-.02,12.35,.43,'white' if i%2==0 else 'bg',radius=True)
    for j,t in enumerate(row):
        text(s,t,col[j],y,[1.96,1.96,1.83,5.65][j],.37,13.5,
             'teal' if i==6 and j==2 else 'ink',j==0)
rect(s,.50,6.19,12.32,.60,'amber_light',radius=True)
text(s,'这些是探索性结果：全部已记录超 2M，且受测 adapter 均为 physics_qualified=false。',.69,6.36,11.96,.36,16,'amber',True)
note(s,'所有数值沿用 18:24:55 CDT 冻结状态。运行状态不等于正式任务有效性或严格预算下的能力结论。',
     ['reports/weekly/2026-09-19/token_audit.json','reports/weekly/2026-09-19/weekly-report.md'])

# 15 — native token use, editable proportional bars
s = base('2M 是目标预算，本批没有实时硬限制', '02 / GPT-6 ASTRA 测试', '原生用量：input + output；缓存输入与推理输出已包含，均不重复相加。')
epmap={e['run_id']:e for e in audit['episodes']}
order=[('cast_coil_codex_001','Coil 001'),('cast_coil_codex_002','Coil 002'),('horseshoe_codex_001','Horseshoe'),('cast_keyhole_codex_001','Keyhole'),('hedgehog_codex_001','Hedgehog'),('lock_picking_codex_001','Lock Picking'),('claws_codex_001','Claws（进行中）')]
barx=2.56;barw=6.28;maximum=70.0
for tick in [0,2,20,40,60,70]:
    x=barx+barw*tick/maximum
    label_w = .15 if tick in [0, 2] else .36
    text(s,str(tick),x-label_w/2,1.98,label_w,.25,10,'red' if tick==2 else 'muted',False,PP_ALIGN.CENTER)
    rect(s,x,2.28,.009,3.46,'red' if tick==2 else 'line')
for i,(rid,label) in enumerate(order):
    y=2.48+i*.46;ep=epmap[rid];n=ep['episode_usage']['total_tokens'];prefix='≥' if not ep['closed_episode'] else ''
    text(s,label,.52,y,1.95,.29,12.8,'ink',True)
    rect(s,barx,y+.025,barw*n/1e6/maximum,.235,'amber' if not ep['closed_episode'] else 'teal',radius=True)
    text(s,f'{prefix}{n:,}',9.15,y-.005,2.0,.30,13.1,'ink',True)
    text(s,f'{prefix}{n/2e6:.2f}×',11.36,y-.005,1.43,.30,13,'red',True)
text(s,'横轴单位：百万 tokens（M）；红线 = 2M',2.57,5.91,8.8,.29,11.8,'muted')
rect(s,.50,6.37,12.31,.43,'red_light',radius=True)
text(s,'7/7 已启动运行均有超预算记录 → 不报告为“严格 2M 下的成功率 / 失败率”。',.69,6.46,11.9,.30,15,'red',True)
note(s,'结束轮次统计含最终答复、排除后续讨论；Claws 为快照下界。Horseshoe 旧快照 2,385,756 更新为完整 episode 2,819,223；Hedgehog 从旧快照 1,934,167 更新为 2,109,472。仅补核原生记录，未修改原 bridge 账本。',
     ['reports/weekly/2026-09-19/token_audit.json'])

# 16 — unstarted tasks and interpretation
s=base('其余 5 项：自主求解未开始', '02 / GPT-6 ASTRA 测试', '已有环境、工程控制器或 adapter 短测，不等于已完成 GPT-6 求解试验。')
for i,(name,body) in enumerate([
 ('Fiber Bundling','已有工程穿环 demo；没有独立 solver 运行记录。'),
 ('Knot Untying','adapter 已接入；CUDA / IPC 等启动问题仍待解决。'),
 ('Buzz Wire 3D','adapter 有 101 步短测，尚无自主求解 episode。'),
 ('Rope Threading','长运行属于工程参考执行，尚未测试 solver。'),
 ('Blanket Flatten','v1–v10 属于工程控制器试验，尚未测试 solver。')]):
    y=1.94+i*.86
    rect(s,.50,y,12.31,.71,'white',radius=True)
    text(s,name,.73,y+.15,2.69,.34,17,'ink',True)
    chip(s,'未开始',3.49,y+.17,1.05,'grey','muted')
    text(s,body,4.91,y+.17,7.57,.41,15.7)
text(s,'正式实验前：同协议参考解通过 → 干净会话 → 外部硬计量 → 冻结条件与重复试验。',.55,6.57,12.0,.35,16,'teal',True)
note(s,'五项任务均明确标为未开始；不将接口 smoke test、工程调参和已有参考 demo 计入 GPT-6 试验。',
     ['reports/weekly/2026-09-19/weekly-report.md'])

# 17 — first reviewer question
s=base('问题一：每个收录变体需要怎样的可解性证据？', '03 / REVIEWER 讨论', '待判断：可复现的成功执行见证是否足够？是否需要更强证明或稳定性重复实验？')
cards=[('几何层','存在满足几何约束的路径','不自动保证机器人能跟随。'),('运动学层','机器人可达、限位与碰撞满足','不自动保证真实动力学执行。'),('执行层','相同评测协议下完整执行通过','支持指定配置下的可解见证。')]
for i,(title,body,foot) in enumerate(cards):
    x=.50+i*4.18
    rect(s,x,2.13,4.0,2.46,'white',radius=True)
    text(s,title,x+.22,2.39,3.56,.41,21,'teal',True)
    text(s,body,x+.22,3.09,3.54,.67,20,'ink',True)
    text(s,foot,x+.22,4.00,3.54,.45,14.5,'muted')
rect(s,.50,5.07,12.32,1.28,'teal_light',radius=True)
text(s,'建议声明范围：最终收录的实例集合',.74,5.27,11.81,.43,23,'teal',True)
text(s,'每个实例绑定私有参考动作与验收记录；不必让 GPT-6 重新盲解每个候选，也不必证明整个连续参数域。',.74,5.86,11.81,.42,17,'ink')
text(s,'一次成功证明指定条件下“存在一条解”；“稳定可解”需要额外的数值或扰动验证。',.55,6.66,12.0,.31,14,'muted')
note(s,'这是待 reviewer 评估的方法建议，并非已完成的全量数据保证。初态、抓持辅助、控制接口、释放与终态判据必须与受测协议一致。',
     ['reports/weekly/2026-09-19/weekly-report.md'])

# 18 — sweep boundary and private witness
s=base('当前阶段：先建立验收闭环，再 sweep 收录集合', '03 / REVIEWER 讨论', '目标是只收录经过验证的实例，同时保留失败候选、拒绝原因与家族分布。')
steps=[('1','已知解与场景\n联合构造'),('2','几何与 IK\n逐例检查'),('3','同一协议\n完整物理执行'),('4','验收通过\n才收录')]
for i,(num,title) in enumerate(steps):
    x=.51+i*3.16
    rect(s,x,2.03,2.84,1.42,'white',radius=True)
    text(s,num,x+.19,2.20,.47,.46,26,'teal',True)
    text(s,title,x+.81,2.25,1.82,.80,18,'ink',True)
    if i<3:text(s,'→',x+2.87,2.46,.29,.38,20,'muted')
for i,(title,body) in enumerate([
 ('先做首批有限 seed','冻结各族成功基准；实现生成、验证、拒绝与有限重采，再扩大规模。'),
 ('声明要与覆盖范围一致','全体可解声明需要逐例证据。只做抽查，就只报告抽查结论。'),
 ('保留评测的独立性','生成端私有参考解；solver 仅使用允许观测。按家族/父模板划分数据。')]):
    x=.51+i*4.18
    rect(s,x,4.03,4.0,2.10,'white',radius=True)
    text(s,title,x+.20,4.28,3.58,.38,19,'ink',True)
    text(s,body,x+.20,4.94,3.58,.95,17,'muted')
text(s,'请 reviewer 确认：逐实例成功见证 + 明示稳定性验证，是否足以支持 paper 的可解性声明？',.54,6.61,12.0,.42,16.6,'teal',True)
note(s,'记录 seed、版本、资产/配置哈希、初态、动作、成功判据与报告。参考控制器失败不等于不可解；有限负对照也不证明唯一解。',
     ['reports/weekly/2026-09-19/weekly-report.md'])

# 19 — actual session testing method
s=base('问题二：直接 GPT-6 session 的测试方法合理吗？', '03 / REVIEWER 讨论', '当前方法评估使用图像、代码与数值工具的长程机器人交互求解。')
rect(s,.50,1.94,3.50,3.57,'white',radius=True)
text(s,'公开输入',.74,2.17,3.0,.43,23,'teal',True)
text(s,'任务目标与规则\n多视角 RGB\n机器人本体状态\n动作接口与相机标定',.74,2.95,2.96,1.96,20)
rect(s,4.20,1.94,4.51,3.57,'teal_light',radius=True)
text(s,'自动执行闭环',4.47,2.17,3.92,.43,23,'teal',True)
text(s,'Observe / 看图\n↓\n规划 / 写动作脚本\n↓\nSubmit / 物理执行 / 再观察',4.47,2.88,3.96,2.29,19,'ink')
rect(s,8.92,1.94,3.90,3.57,'white',radius=True)
text(s,'规则与结果',9.17,2.17,3.39,.43,23,'teal',True)
text(s,'单一连续 episode\n不重置、不回滚、不改环境\n不读已有解法 / 私有状态\nEvaluator 判定成功',9.17,2.95,3.34,1.96,17.5)
text(s,'本地文件队列或 HTTP bridge；manual 仅指计量模式，动作不需要人逐段批准。',.55,5.85,12.0,.44,17,'muted')
rect(s,.50,6.44,12.32,.42,'amber_light',radius=True)
text(s,'当前 2M 没有外部实时硬限制；工具与资料边界主要依赖提示词约定。',.70,6.52,11.93,.31,15.6,'amber',True)
note(s,'允许终端、通用库、图像分析和自行编写的 IK/规划；联网由配置固定。代码预检查拒绝不推进仿真，网络结果不明用同 request ID 重试。ready 不表示成功。',
     ['docs/develop/modules/agent_bridge_zh.md','docs/develop/modules/agent_solver_prompt_zh.md',
      'docs/develop/modules/agent_test_demos_zh.md'])

# 20 — reviewer decisions and constraints, no new experimental claims
s=base('正式 benchmark 前，需要 reviewer 评估的四项条件', '03 / REVIEWER 讨论', '现有这批适合作为探索性 pilot；更强的论文结论需要补齐以下条件。')
criteria=[
 ('01','信息与会话独立','新会话 + 干净任务包 / 容器；固定允许的机器人信息、标定和工具。','现状：部分任务沿用有旧上下文的同一聊天。'),
 ('02','真实 2M 硬预算','外部采集原生 usage，调用前预留额度，动作前核验；明确报告计量范围。','现状：7 次运行均已记录超预算，只有事后或临时核对。'),
 ('03','同协议可解见证','参考控制器通过同一初态、辅助抓持、动作接口、终态与失败规则。','现状：工程 demo 与 solver 初态及抓持条件存在差异。'),
 ('04','冻结条件与重复试验','固定模型、工具、观测、网络、实例和 evaluator；分层报告失败类型。','现状：少量试验，版本、抓持条件和物理后端不完全一致。')]
for i,(num,title,body,now) in enumerate(criteria):
    x=.50+(i%2)*6.29;y=1.98+(i//2)*2.35
    rect(s,x,y,6.05,2.09,'white',radius=True)
    text(s,num,x+.20,y+.16,.56,.43,26,'teal',True)
    text(s,title,x+.91,y+.20,4.90,.38,21,'ink',True)
    text(s,body,x+.22,y+.81,5.58,.65,16.5)
    text(s,now,x+.22,y+1.57,5.58,.39,12.5,'muted')
linked_text(s,'完整周报与 73 个多视角 MP4 ↗',DATA/'weekly-report.md',.56,6.88,9.3,.24,10.6)
note(s,'请 reviewer 确认：1）最终收录实例的可解见证及稳定性最低标准；2）直接 coding-agent session 的信息/工具范围及正式 2M 评测控制条件。',
     ['reports/weekly/2026-09-19/weekly-report.md', 'reports/weekly/2026-09-19/token_audit.json'])

# Keep all text/charts editable; embed the complete primary multiview videos.
assert len(prs.slides) == 20
assert len(embedded_movies) == 39
prs.save(DECK)
manifest={'pptx':str(DECK),'slides':slide_records,'hyperlinks':links,
          'linked_mp4_unique':len({x['source'] for x in links if x['source'].endswith('.mp4')}),
          'video_embedding':'39 complete H.264 MP4 movies embedded in the PPTX; click in Slide Show to play',
          'embedded_movies':embedded_movies,
          'claws_primary_variant':'A0',
          'source_report':str(DATA/'weekly-report.md'),
          'snapshot_engineering':'2026-09-19 18:27 CDT','snapshot_agent':'2026-09-19 18:24:55 CDT'}
(OUT/'TopoGym_Weekly_2026-09-19.manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(f'Saved {DECK}; {len(prs.slides)} slides, 39 embedded movies, {manifest["linked_mp4_unique"]} supplemental MP4 links')

# Render this separate, poster-only copy for a small static PDF. LibreOffice's
# ordinary PDF export otherwise embeds all 56 MB of videos into PDF annotations.
for slide in prs.slides:
    timing = slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing')
    if timing is not None:
        slide._element.remove(timing)
    for shape in slide.shapes:
        if shape.name.startswith('DEMO '):
            for node in shape._element.xpath('.//a:videoFile | .//p:nvPr/p:extLst | .//a:hlinkClick'):
                node.getparent().remove(node)
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith(('/media', '/video')):
            slide.part.drop_rel(rel.rId)
prs.save(ASSETS/'static_preview.pptx')
